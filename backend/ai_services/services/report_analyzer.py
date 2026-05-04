"""Extract text from a report file and produce summary/weaknesses/suggestions.

If ANTHROPIC_API_KEY or OPENAI_API_KEY is set, this would call the LLM.
Otherwise it produces a deterministic heuristic analysis so the system runs
end-to-end without external dependencies.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from django.conf import settings


@dataclass
class ReportInsights:
    summary: str
    weaknesses: list[str]
    suggestions: list[str]
    raw: dict


def _extract_text(path: str) -> str:
    p = Path(path)
    if not p.exists():
        return ""
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(str(p))
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception:
            return ""
    if suffix in (".docx", ".doc"):
        try:
            import docx  # python-docx
            d = docx.Document(str(p))
            return "\n".join(par.text for par in d.paragraphs)
        except Exception:
            return ""
    if suffix in (".txt", ".md"):
        return p.read_text(errors="ignore")
    return ""


def _heuristic_analysis(text: str) -> ReportInsights:
    text = text.strip()
    words = re.findall(r"\b\w+\b", text)
    word_count = len(words)
    sentences = re.split(r"[.!?]+", text)
    sentence_count = max(1, len([s for s in sentences if s.strip()]))
    avg_sentence = word_count / sentence_count

    weaknesses: list[str] = []
    suggestions: list[str] = []

    if word_count < 800:
        weaknesses.append(f"Report is short ({word_count} words). Expand the analysis sections.")
        suggestions.append("Add a methodology section and discuss tradeoffs of design decisions.")
    if avg_sentence > 28:
        weaknesses.append("Average sentence length is high — readability suffers.")
        suggestions.append("Break long sentences and prefer active voice.")
    if "TODO" in text or "TBD" in text:
        weaknesses.append("Placeholder markers (TODO/TBD) remain in the text.")
        suggestions.append("Remove placeholders before final submission.")

    headings = re.findall(r"^(#+\s+.+|[0-9]+\.\s+[A-Z].{3,80})$", text, flags=re.MULTILINE)
    if len(headings) < 3:
        weaknesses.append("Few section headings detected — structure is unclear.")
        suggestions.append("Add headings for Introduction, Approach, Results, and Conclusion.")

    if "test" not in text.lower() and "evaluation" not in text.lower():
        weaknesses.append("No mention of testing or evaluation.")
        suggestions.append("Document how the system was tested and what metrics were used.")

    summary = (
        f"Report contains ~{word_count} words across ~{sentence_count} sentences "
        f"(avg {avg_sentence:.1f} words/sentence) with {len(headings)} headings detected."
    )
    return ReportInsights(
        summary=summary,
        weaknesses=weaknesses or ["No major issues detected by heuristic scan."],
        suggestions=suggestions or ["Consider a peer read-through for clarity."],
        raw={"word_count": word_count, "sentences": sentence_count, "headings": len(headings)},
    )


def analyze(report_path: str) -> ReportInsights:
    text = _extract_text(report_path)
    if not text:
        return ReportInsights(
            summary="Could not extract text from report.",
            weaknesses=["Report file unreadable or empty."],
            suggestions=["Re-upload as PDF, DOCX, TXT, or MD."],
            raw={},
        )
    # LLM hook would go here. We keep the deterministic path for now.
    if settings.OPENAI_API_KEY or settings.ANTHROPIC_API_KEY:
        # Placeholder: in production, call provider; falling back to heuristic for repeatability.
        pass
    return _heuristic_analysis(text)
